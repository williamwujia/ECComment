from __future__ import annotations

import json
import os
import shutil
import sqlite3
import unicodedata
import uuid
from collections import Counter, defaultdict
from contextlib import closing
from datetime import date, datetime
from pathlib import Path, PurePosixPath

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


REVIEW_EXPORT_COLUMNS = [
    "platform", "product_id", "product_title", "local_review_id",
    "platform_review_id", "review_time", "append_time",
    "reviewer_display_name", "sku", "review_text", "append_text",
    "new_review_image_count", "new_append_image_count",
    "new_total_image_count", "export_review_image_count",
    "export_append_image_count", "export_total_image_count",
    "review_folder", "source_file_name",
]
IMAGE_EXPORT_COLUMNS = [
    "platform", "product_id", "product_title", "local_review_id",
    "image_stage", "image_index", "is_new_in_batch", "review_time", "review_text",
    "append_text", "查看图片", "relative_path", "file_name",
    "export_status", "source_file_name",
]


def _load_json(value: str | None) -> dict:
    loaded = json.loads(value or "null")
    return loaded if isinstance(loaded, dict) else {}


def _batch_snapshots(db: sqlite3.Connection, batch_id: str) -> tuple[dict, dict, dict, list[dict]]:
    batch = db.execute("SELECT * FROM batches WHERE batch_id=?", (batch_id,)).fetchone()
    if batch is None:
        raise KeyError("找不到该批次")
    if batch["status"] != "completed":
        raise ValueError("只有未撤销的已完成批次可以导出图片包")
    products: dict[str, dict] = {}
    reviews: dict[str, dict] = {}
    image_changes: list[dict] = []
    changes = db.execute(
        "SELECT * FROM batch_changes WHERE batch_id=? ORDER BY change_id",
        (batch_id,),
    ).fetchall()
    for change in changes:
        state = _load_json(change["new_state"])
        if change["entity_type"] == "product" and state:
            products[change["entity_id"]] = state
        elif change["entity_type"] == "review" and state:
            state["_created_in_batch"] = (
                change["operation"] == "create"
                or reviews.get(change["entity_id"], {}).get("_created_in_batch", False)
            )
            state["_operation"] = change["operation"]
            reviews[change["entity_id"]] = state
        elif change["entity_type"] == "image" and change["operation"] == "create" and state:
            image_changes.append(state)
    return dict(batch), products, reviews, image_changes


def _unique_destination(export_root: Path, batch_id: str) -> Path:
    base_name = f"{date.today().isoformat()}_评论图片导出_{batch_id}"
    candidate = export_root / base_name
    suffix = 2
    while candidate.exists():
        candidate = export_root / f"{base_name}_{suffix}"
        suffix += 1
    return candidate


def _write_workbook(
    path: Path,
    summary_rows: list[tuple[str, object]],
    product_rows: list[dict],
    review_rows: list[dict],
    image_rows: list[dict],
) -> None:
    def display_width(value: object) -> int:
        return sum(
            2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
            for character in str(value or "")
        )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["指标", "值"])
    for label, value in summary_rows:
        summary.append([label, value])
    summary.append([])
    summary.append(["商品明细"])
    product_columns = [
        "platform", "product_id", "product_title", "new_review_count",
        "new_image_count", "export_image_count",
    ]
    summary.append(product_columns)
    for row in product_rows:
        summary.append([row.get(column, "") for column in product_columns])

    reviews_sheet = workbook.create_sheet("reviews")
    reviews_sheet.append(REVIEW_EXPORT_COLUMNS)
    for row in review_rows:
        reviews_sheet.append([row.get(column, "") for column in REVIEW_EXPORT_COLUMNS])

    images_sheet = workbook.create_sheet("images")
    images_sheet.append(IMAGE_EXPORT_COLUMNS)
    link_column = IMAGE_EXPORT_COLUMNS.index("查看图片") + 1
    for row in image_rows:
        images_sheet.append([row.get(column, "") for column in IMAGE_EXPORT_COLUMNS])
        cell = images_sheet.cell(images_sheet.max_row, link_column)
        if row.get("export_status") == "success":
            cell.value = "查看图片"
            cell.hyperlink = row["relative_path"]
            cell.style = "Hyperlink"
        else:
            cell.value = "导出失败"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (summary, reviews_sheet, images_sheet):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column in sheet.columns:
            letter = column[0].column_letter
            width = max((display_width(cell.value) for cell in column[:100]), default=10)
            sheet.column_dimensions[letter].width = min(max(width + 2, 12), 48)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[1].height = 24
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 42
    summary["A15"].font = Font(bold=True, color="1F4E78")
    for cell in summary[16]:
        cell.fill = header_fill
        cell.font = header_font

    text_columns = {"product_id", "local_review_id", "platform_review_id", "sku"}
    for sheet in (reviews_sheet, images_sheet):
        headers = {cell.value: cell.column for cell in sheet[1]}
        for header in text_columns.intersection(headers):
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row_index, headers[header]).number_format = "@"
    for row_index in range(17, summary.max_row + 1):
        summary.cell(row_index, 2).number_format = "@"
    workbook.save(path)


def export_batch_package(
    *, asset_root: Path, db_path: Path, batch_id: str, export_root: str | Path
) -> dict:
    export_base = Path(export_root).expanduser().resolve()
    export_base.mkdir(parents=True, exist_ok=True)
    destination = _unique_destination(export_base, batch_id)
    staging = export_base / f".{destination.name}.tmp-{uuid.uuid4().hex}"
    images_root = staging / "images"
    staging.mkdir(parents=True)
    images_root.mkdir()

    with closing(sqlite3.connect(db_path)) as db:
        db.row_factory = sqlite3.Row
        batch, products, reviews, image_changes = _batch_snapshots(db, batch_id)
        file_rows = [dict(row) for row in db.execute(
            "SELECT * FROM batch_files WHERE batch_id=?", (batch_id,)
        ).fetchall()]
        identified_images = sum(int(row.get("images_found") or 0) for row in file_rows)
        observed_rows = [dict(row) for row in db.execute(
            """SELECT i.*, o.was_new, o.source_file_name AS observed_source_file_name
               FROM batch_image_observations o
               JOIN images i ON i.image_id=o.image_id
               WHERE o.batch_id=? ORDER BY i.local_review_id,i.image_stage,i.image_index""",
            (batch_id,),
        ).fetchall()]
        legacy_warning = ""
        if observed_rows:
            export_images = observed_rows
        elif identified_images > len(image_changes):
            product_keys = json.loads(batch.get("products_affected") or "[]")
            placeholders = ",".join("?" for _ in product_keys)
            changed_ids = {str(image.get("image_id", "")) for image in image_changes}
            export_images = []
            if product_keys:
                export_images = [dict(row) for row in db.execute(
                    f"""SELECT i.* FROM images i
                        JOIN reviews r ON r.local_review_id=i.local_review_id
                        WHERE r.product_key IN ({placeholders})
                        ORDER BY r.product_key,i.local_review_id,i.image_stage,i.image_index""",
                    product_keys,
                ).fetchall()]
                for image in export_images:
                    image["was_new"] = int(image.get("image_id", "") in changed_ids)
                    image["observed_source_file_name"] = image.get("source_file_name", "")
            legacy_warning = "旧批次缺少逐图识别关系，已按受影响商品的当前图片资产生成非空目录"
        else:
            export_images = []
            for snapshot in image_changes:
                image = dict(snapshot)
                image["was_new"] = 1
                image["observed_source_file_name"] = image.get("source_file_name", "")
                export_images.append(image)

        missing_review_ids = sorted({
            str(image.get("local_review_id", "")) for image in export_images
            if str(image.get("local_review_id", "")) not in reviews
        })
        if missing_review_ids:
            placeholders = ",".join("?" for _ in missing_review_ids)
            for row in db.execute(
                f"SELECT * FROM reviews WHERE local_review_id IN ({placeholders})",
                missing_review_ids,
            ).fetchall():
                reviews[row["local_review_id"]] = dict(row)
        missing_product_keys = sorted({
            str(review.get("product_key", "")) for review in reviews.values()
            if str(review.get("product_key", "")) not in products
        })
        if missing_product_keys:
            placeholders = ",".join("?" for _ in missing_product_keys)
            for row in db.execute(
                f"SELECT * FROM products WHERE product_key IN ({placeholders})",
                missing_product_keys,
            ).fetchall():
                products[row["product_key"]] = dict(row)

    created_review_ids = {
        review_id for review_id, review in reviews.items()
        if review.get("_created_in_batch")
    }
    new_image_counts: dict[str, Counter] = defaultdict(Counter)
    export_image_counts: dict[str, Counter] = defaultdict(Counter)
    image_rows: list[dict] = []
    copy_warnings: list[str] = [legacy_warning] if legacy_warning else []
    copied_count = 0

    try:
        for image in export_images:
            local_id = str(image.get("local_review_id", ""))
            review = reviews.get(local_id, {})
            product_key = str(review.get("product_key", ""))
            product = products.get(product_key, {})
            stage = str(image.get("image_stage", ""))
            was_new = bool(image.get("was_new"))
            export_image_counts[local_id][stage] += 1
            if was_new:
                new_image_counts[local_id][stage] += 1
            relative_path = (
                PurePosixPath("images") / product_key / local_id / str(image.get("file_name", ""))
            ).as_posix()
            source_relative = str(image.get("local_path", ""))
            source = asset_root / Path(source_relative) if source_relative else None
            target = staging / Path(relative_path)
            export_status = "failed"
            if source is not None and source.is_file() and image.get("extract_status") == "extracted":
                try:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, target)
                    export_status = "success"
                    copied_count += 1
                except OSError as exc:
                    copy_warnings.append(f"{image.get('image_id', '')}：{exc}")
            else:
                copy_warnings.append(f"{image.get('image_id', '')}：长期资产图片不存在或未成功提取")
            image_rows.append({
                "platform": product.get("platform", ""),
                "product_id": product.get("product_id", ""),
                "product_title": product.get("product_title", ""),
                "local_review_id": local_id,
                "image_stage": stage,
                "image_index": image.get("image_index", ""),
                "is_new_in_batch": was_new,
                "review_time": review.get("review_time", ""),
                "review_text": review.get("review_text", ""),
                "append_text": review.get("append_text", ""),
                "查看图片": "查看图片" if export_status == "success" else "导出失败",
                "relative_path": relative_path,
                "file_name": image.get("file_name", ""),
                "export_status": export_status,
                "source_file_name": image.get("observed_source_file_name", image.get("source_file_name", "")),
            })

        review_rows: list[dict] = []
        for local_id in sorted(export_image_counts):
            review = reviews.get(local_id, {})
            product_key = str(review.get("product_key", ""))
            product = products.get(product_key, {})
            new_counts = new_image_counts[local_id]
            export_counts = export_image_counts[local_id]
            review_rows.append({
                "platform": product.get("platform", ""),
                "product_id": product.get("product_id", ""),
                "product_title": product.get("product_title", ""),
                "local_review_id": local_id,
                "platform_review_id": review.get("platform_review_id", "") or "",
                "review_time": review.get("review_time", ""),
                "append_time": review.get("append_time", ""),
                "reviewer_display_name": review.get("reviewer_display_name", ""),
                "sku": review.get("sku", ""),
                "review_text": review.get("review_text", ""),
                "append_text": review.get("append_text", ""),
                "new_review_image_count": new_counts["review"],
                "new_append_image_count": new_counts["append"],
                "new_total_image_count": sum(new_counts.values()),
                "export_review_image_count": export_counts["review"],
                "export_append_image_count": export_counts["append"],
                "export_total_image_count": sum(export_counts.values()),
                "review_folder": (PurePosixPath("images") / product_key / local_id).as_posix(),
                "source_file_name": review.get("source_file_name", ""),
            })

        product_image_counts = Counter()
        product_export_counts = Counter()
        product_review_counts = Counter()
        for image in export_images:
            review = reviews.get(str(image.get("local_review_id", "")), {})
            product_key = str(review.get("product_key", ""))
            product_export_counts[product_key] += 1
            if image.get("was_new"):
                product_image_counts[product_key] += 1
        for local_id in created_review_ids:
            product_review_counts[str(reviews[local_id].get("product_key", ""))] += 1
        product_rows = [
            {
                "platform": product.get("platform", ""),
                "product_id": product.get("product_id", ""),
                "product_title": product.get("product_title", ""),
                "new_review_count": product_review_counts[product_key],
                "new_image_count": product_image_counts[product_key],
                "export_image_count": product_export_counts[product_key],
            }
            for product_key, product in sorted(products.items())
        ]

        identified_reviews = sum(int(row.get("reviews_found") or 0) for row in file_rows)
        new_images = sum(bool(image.get("was_new")) for image in export_images)
        historical_images = len(export_images) - new_images
        warning_count = sum(int(row.get("warning_count") or 0) for row in file_rows) + len(copy_warnings)
        summary_rows = [
            ("batch_id", batch_id),
            ("导出时间", datetime.now().astimezone().isoformat(timespec="seconds")),
            ("商品数", len(products)),
            ("本批次识别评论数", identified_reviews),
            ("本批次新增评论数", int(batch.get("reviews_created") or 0)),
            ("有导出图片的评论数", len(review_rows)),
            ("识别图片数", identified_images),
            ("本次导出历史图片数", historical_images),
            ("本次新增原评图片数", sum(row["image_stage"] == "review" and row["is_new_in_batch"] for row in image_rows)),
            ("本次新增追评图片数", sum(row["image_stage"] == "append" and row["is_new_in_batch"] for row in image_rows)),
            ("本次实际导出图片数", copied_count),
            ("警告数", warning_count),
        ]
        workbook_path = staging / "图片目录.xlsx"
        _write_workbook(workbook_path, summary_rows, product_rows, review_rows, image_rows)

        actual_files = sum(1 for path in images_root.rglob("*") if path.is_file()) if images_root.exists() else 0
        broken_links = 0
        for row in image_rows:
            relative = PurePosixPath(str(row["relative_path"]))
            safe_relative = not relative.is_absolute() and ".." not in relative.parts
            if row["export_status"] == "success" and (
                not safe_relative or not (staging / Path(*relative.parts)).is_file()
            ):
                broken_links += 1
        workbook = load_workbook(workbook_path, read_only=False)
        try:
            images_sheet = workbook["images"]
            excel_image_rows = max(images_sheet.max_row - 1, 0)
            link_column = IMAGE_EXPORT_COLUMNS.index("查看图片") + 1
            for row_index, row in enumerate(image_rows, 2):
                cell = images_sheet.cell(row_index, link_column)
                if row["export_status"] == "success":
                    target = cell.hyperlink.target if cell.hyperlink else ""
                    if target != row["relative_path"] or Path(target).is_absolute() or ".." in PurePosixPath(target).parts:
                        broken_links += 1
        finally:
            workbook.close()
        failed_count = len(image_rows) - copied_count
        inconsistent = actual_files != copied_count or excel_image_rows != len(image_rows) or broken_links > 0
        status = "导出完成，有警告" if failed_count or inconsistent or copy_warnings else "导出完成"
        os.replace(staging, destination)
        return {
            "batch_id": batch_id,
            "status": status,
            "export_directory": str(destination),
            "workbook_path": str(destination / "图片目录.xlsx"),
            "identified_image_count": identified_images,
            "historical_image_count": historical_images,
            "new_image_count": new_images,
            "export_candidate_count": len(export_images),
            "exported_image_count": copied_count,
            "failed_image_count": failed_count,
            "excel_image_row_count": excel_image_rows,
            "actual_file_count": actual_files,
            "broken_link_count": broken_links,
            "warnings": copy_warnings,
        }
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
