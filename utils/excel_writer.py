from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_COLUMNS = {
    "ai_candidates": [
        "content_id", "source_file", "platform", "product_title", "shop_name",
        "product_url", "product_id", "record_type", "content_role",
        "content_text_clean", "parent_text_clean", "review_time", "review_date",
        "user_status", "sku",
        "pre_purchase_decision", "pre_purchase_terms", "pre_purchase_reason",
        "ai_candidate", "ai_candidate_gate_reason",
        "ai_influence_level", "ai_evidence_type", "ai_level_reason",
        "ai_source_terms", "ai_action_terms", "purchase_context_terms",
        "decision_context_terms", "research_terms", "algorithm_terms",
        "source_type", "matched_keywords", "matched_sentence",
        "match_score", "evidence_level", "evidence_reason", "excluded_by_rule",
        "exclude_reason", "content_hash",
    ],
    "content_all": [
        "content_id", "record_type", "source_file", "platform", "product_title",
        "shop_name", "product_url", "product_id", "content_role",
        "content_text_raw", "content_text_clean", "parent_text_clean",
        "review_time", "review_date", "sku",
        "user_name_masked", "user_status", "content_order", "content_hash",
        "pre_purchase_decision", "pre_purchase_terms", "pre_purchase_reason",
        "ai_candidate", "ai_candidate_gate_reason", "ai_related",
        "ai_influence_level", "ai_evidence_type", "ai_level_reason",
        "ai_source_hit", "ai_source_terms", "ai_action_hit", "ai_action_terms",
        "purchase_context_hit", "purchase_context_terms", "decision_context_hit",
        "decision_context_terms", "research_terms", "algorithm_terms",
        "source_type", "matched_keywords", "matched_sentence",
        "match_score", "evidence_level", "evidence_reason", "excluded_by_rule",
        "exclude_reason", "extract_confidence",
    ],
    "qa_pairs_raw": [
        "source_file", "platform", "product_title", "shop_name", "product_url",
        "product_id", "qa_order", "question_text_raw", "question_text_clean",
        "answer_text_raw", "answer_text_clean", "answer_user_status",
        "answer_count", "question_tags", "is_buyer_answer", "qa_hash",
        "extract_confidence",
    ],
    "reviews_raw": [
        "source_file", "platform", "product_title", "shop_name", "product_url",
        "product_id", "review_order", "review_text_raw", "review_text_clean",
        "review_time", "review_date", "rating", "sku", "user_name_masked",
        "is_followup",
        "followup_text", "merchant_reply", "image_count", "video_count",
        "like_count", "review_hash", "extract_confidence",
    ],
    "summary_by_product": [
        "platform", "product_title", "shop_name", "product_id",
        "source_file_count", "review_count", "qa_question_count",
        "qa_answer_count", "content_count", "ai_candidate_count",
        "a_level_count", "b_level_count", "c_level_count", "d_level_count",
        "a_b_count", "a_b_c_count", "pre_purchase_decision_count",
        "explicit_ai_rate", "generic_ai_rate", "inferred_ai_rate",
        "ai_candidate_rate",
        "top_ai_source_terms", "top_ai_action_terms", "top_examples",
    ],
    "summary_by_keyword": [
        "keyword_type", "keyword", "hit_count", "example_1", "example_2",
        "example_3",
    ],
    "summary_by_month": [
        "review_month", "review_count", "pre_purchase_decision_count", "ai_candidate_count",
        "a_level_count", "b_level_count", "c_level_count",
        "ai_candidate_rate", "top_ai_source_terms",
    ],
    "pre_purchase_details": [
        "content_id", "record_type", "source_file", "platform", "product_title",
        "content_role", "content_text_clean", "parent_text_clean",
        "pre_purchase_terms", "pre_purchase_reason", "research_terms",
        "source_type", "ai_candidate", "ai_influence_level",
        "ai_evidence_type", "content_hash",
    ],
    "sentiment_run_summary": [
        "enabled", "status", "message", "target_review_count", "rule_count",
        "llm_fast_count", "detail_count", "failed_count", "sentiment_limit",
        "detail_limit", "sentiment_batch_size", "detail_batch_size",
        "dry_run", "elapsed_seconds",
    ],
    "sentiment_fast": [
        "content_id", "content_hash", "source_file", "platform", "product_title",
        "sku", "content_text_clean", "sentiment_label", "sentiment_label_cn",
        "sentiment_score", "praise_code", "praise_cn", "complaint_code",
        "complaint_cn", "sentiment_confidence", "sentiment_source",
    ],
    "sentiment_detail": [
        "content_id", "content_hash", "source_file", "platform", "product_title",
        "sku", "content_text_clean", "sentiment_label", "sentiment_label_cn",
        "sentiment_score", "praise_code", "praise_cn", "complaint_code",
        "complaint_cn", "sentiment_confidence", "sentiment_source",
        "sentiment_evidence", "sentiment_reason", "geo_value",
    ],
    "sentiment_failures": [
        "stage", "content_id", "content_hash", "content_text_clean",
        "error_type", "error_message", "raw_response",
    ],
    "sentiment_timings": [
        "scope", "batch_size", "attempt", "status", "error_type",
        "error_message", "total_seconds", "response_read_seconds",
        "prompt_tokens", "completion_tokens", "total_tokens", "max_tokens",
        "finish_reason", "content_ids",
    ],
    "sentiment_summary": [
        "metric", "label", "value",
    ],
    "llm_run_summary": [
        "enabled", "status", "message", "provider", "model", "prompt_version",
        "target_review_count", "total_count", "success_count", "failed_count",
        "skipped_count", "config_path", "api_key_file_override", "dry_run",
        "batch_size", "batch_size_source", "target_content_count",
        "target_total_seconds", "timeout_seconds", "retry_delays",
        "llm_elapsed_seconds",
    ],
    "llm_review_analysis": [
        "content_id", "comment_id", "source_file", "platform", "product_title",
        "product_id", "sku", "comment_text", "overall_sentiment", "main_target",
        "purchase_evidence_level", "is_purchase_decision_evidence",
        "purchase_evidence_reason", "praise_count", "complaint_count",
        "noise_flags", "one_sentence_summary", "model_name", "prompt_version",
        "input_hash", "usage_json", "raw_llm_json",
    ],
    "llm_praise_items": [
        "content_id", "comment_id", "source_file", "platform", "product_title",
        "product_id", "sku", "comment_text", "item_index", "target", "aspect",
        "praise_family", "praise_method", "scene", "evidence_quote",
        "evidence_strength", "business_value", "notes",
    ],
    "llm_complaint_items": [
        "content_id", "comment_id", "source_file", "platform", "product_title",
        "product_id", "sku", "comment_text", "item_index", "target", "aspect",
        "complaint_family", "complaint_method", "scene", "evidence_quote",
        "evidence_strength", "severity", "fixability", "notes",
    ],
    "llm_failures": [
        "content_id", "comment_id", "source_file", "product_title",
        "comment_text", "model_name", "prompt_version", "input_hash",
        "error_type", "error_message", "raw_response",
    ],
    "llm_request_timings": [
        "request_index", "scope", "model_name", "prompt_version", "batch_size",
        "attempt", "status", "error_type", "error_message", "total_seconds",
        "payload_build_seconds", "http_open_seconds", "response_read_seconds",
        "json_parse_seconds", "request_bytes", "response_bytes", "prompt_tokens",
        "completion_tokens", "total_tokens", "max_tokens", "timeout_seconds",
        "finish_reason", "content_ids",
    ],
    "errors": [
        "source_file", "platform", "error_stage", "error_type",
        "error_message", "traceback",
    ],
    "debug_samples": [
        "source_file", "platform", "debug_type", "raw_block_text",
        "possible_question", "possible_answer", "reason",
    ],
}


def _frame(name: str, rows: list[dict]) -> pd.DataFrame:
    columns = SHEET_COLUMNS[name]
    frame = pd.DataFrame(rows)
    for column in columns:
        if column not in frame.columns:
            frame[column] = None
    return frame[columns]


def write_excel(output_path: str, sheets: dict) -> None:
    """Write the required workbook with stable columns and readable styling."""
    path = Path(output_path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in SHEET_COLUMNS:
            _frame(name, sheets.get(name, [])).to_excel(
                writer, sheet_name=name, index=False
            )

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for column_index, cells in enumerate(sheet.columns, 1):
                values = [
                    len(str(cell.value)) for cell in list(cells)[:200] if cell.value
                ]
                width = min(max(values or [8]) + 2, 45)
                sheet.column_dimensions[get_column_letter(column_index)].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
