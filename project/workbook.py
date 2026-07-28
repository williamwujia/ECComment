from __future__ import annotations

import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from project.schema import COMPATIBILITY_SHEETS, SHEET_COLUMNS
from utils.excel_sanitizer import sanitize_frame_for_excel


def empty_frame(sheet_name: str) -> pd.DataFrame:
    return pd.DataFrame(columns=SHEET_COLUMNS.get(sheet_name, []))


def load_sheets(workbook_path: str | Path) -> dict[str, pd.DataFrame]:
    path = Path(workbook_path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"项目工作簿不存在：{path}")
    sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    for name in SHEET_COLUMNS:
        if name not in sheets:
            sheets[name] = empty_frame(name)
        else:
            sheets[name] = normalize_frame(name, sheets[name])
    for name in COMPATIBILITY_SHEETS:
        sheets.setdefault(name, pd.DataFrame())
    return sheets


def normalize_frame(name: str, frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in SHEET_COLUMNS.get(name, []):
        if column not in result.columns:
            result[column] = None
    columns = SHEET_COLUMNS.get(name)
    return result[columns] if columns else result


def validate_sheets(sheets: dict[str, pd.DataFrame]) -> None:
    missing = [name for name in SHEET_COLUMNS if name not in sheets]
    if missing:
        raise ValueError(f"工作簿缺少管理 Sheet：{', '.join(missing)}")
    projects = sheets["project_info"]
    if projects.empty or projects["project_id"].fillna("").astype(str).str.strip().eq("").any():
        raise ValueError("project_info 必须至少包含一个有效项目")
    products = sheets["products"]
    if not products.empty:
        duplicated = products.duplicated(["project_id", "item_key"], keep=False)
        if duplicated.any():
            raise ValueError("products 中存在重复的 (project_id, item_key)")


def write_sheets_atomic(
    workbook_path: str | Path,
    sheets: dict[str, pd.DataFrame],
    *,
    create_backup: bool = True,
) -> Path | None:
    """Write all sheets to a sibling temp file, validate, back up, then replace."""
    path = Path(workbook_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    for name in SHEET_COLUMNS:
        sheets[name] = normalize_frame(name, sheets.get(name, empty_frame(name)))
    validate_sheets(sheets)

    handle = tempfile.NamedTemporaryFile(
        prefix=f".{path.stem}_",
        suffix=".xlsx",
        dir=path.parent,
        delete=False,
    )
    temp_path = Path(handle.name)
    handle.close()
    backup_path: Path | None = None
    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            ordered = list(SHEET_COLUMNS) + [
                name for name in sheets if name not in SHEET_COLUMNS
            ]
            for name in ordered:
                frame = sanitize_frame_for_excel(sheets.get(name, pd.DataFrame()))
                frame.to_excel(writer, sheet_name=name[:31], index=False)
            _style_workbook(writer.book)

        check = load_workbook(temp_path, read_only=True, data_only=True)
        required = set(SHEET_COLUMNS)
        if not required.issubset(check.sheetnames):
            raise ValueError("临时工作簿校验失败：管理 Sheet 不完整")
        check.close()

        if path.exists() and create_backup:
            backup_dir = path.parent / "backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup_path = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
            shutil.copy2(path, backup_path)
        os.replace(temp_path, path)
        return backup_path
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _style_workbook(workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for column_index, cells in enumerate(sheet.columns, 1):
            sample = list(cells)[:200]
            width = min(max([len(str(cell.value)) for cell in sample if cell.value] or [8]) + 2, 45)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
